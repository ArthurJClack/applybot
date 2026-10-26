"""Render resumes/cover letters to PDF (ATS-friendly single column), DOCX, and Markdown/plain text."""
from __future__ import annotations

import os
import re
from datetime import date
from pathlib import Path

from fpdf import FPDF

_UNICODE_FONTS = [
    (r"C:\Windows\Fonts\calibri.ttf", r"C:\Windows\Fonts\calibrib.ttf", r"C:\Windows\Fonts\calibrii.ttf"),
    (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf", r"C:\Windows\Fonts\ariali.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf"),
    ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf", "/Library/Fonts/Arial Italic.ttf"),
]
_LATIN1_MAP = str.maketrans({"\u2013": "-", "\u2014": "-", "\u2018": "'", "\u2019": "'", "\u201c": '"',
                             "\u201d": '"', "\u2022": "-", "\u2026": "...", "\u00a0": " ", "\u2192": "->"})


def _contact_line(c: dict) -> str:
    keys = ["email", "phone", "location", "linkedin", "github", "website"]
    return "  |  ".join(str(c[k]) for k in keys if c.get(k))


def _entry_header(section: str, e: dict) -> tuple[str, str, str]:
    """Returns (bold left, right/dates, italic second line)."""
    if section == "education":
        sub = ", ".join(x for x in [e.get("degree"), f"GPA: {e['gpa']}" if e.get("gpa") else ""] if x)
        return e.get("school", ""), e.get("dates", ""), " - ".join(x for x in [sub, e.get("location")] if x)
    if section == "projects":
        return e.get("name", ""), e.get("dates", ""), " | ".join(x for x in [e.get("tech"), e.get("link")] if x)
    return e.get("title", ""), e.get("dates", ""), " - ".join(x for x in [e.get("company"), e.get("location")] if x)


SECTIONS = [("education", "Education"), ("experience", "Experience"), ("projects", "Projects"),
            ("leadership", "Leadership & Activities")]


def _is_student(education: list[dict]) -> bool:
    """Students/recent grads list education first."""
    this_year = date.today().year
    for e in education:
        d = str(e.get("dates", "")).lower()
        years = [int(y) for y in re.findall(r"(?:19|20)\d{2}", d)]
        if "expected" in d or "present" in d or (years and max(years) >= this_year - 1):
            return True
    return False


def section_order(resume: dict) -> list[tuple[str, str]]:
    """Education first for students/recent grads; experience first for everyone else."""
    if resume.get("experience") and not _is_student(resume.get("education") or []):
        return [SECTIONS[1], SECTIONS[0], *SECTIONS[2:]]
    return SECTIONS[:]


class _PDF(FPDF):
    def __init__(self):
        super().__init__(format="Letter")
        self.family = "Helvetica"
        self.unicode = False
        for regular, bold, italic in _UNICODE_FONTS:
            if all(os.path.exists(p) for p in (regular, bold, italic)):
                self.add_font("Body", "", regular)
                self.add_font("Body", "B", bold)
                self.add_font("Body", "I", italic)
                self.family, self.unicode = "Body", True
                break

    def t(self, s) -> str:
        s = str(s or "")
        return s if self.unicode else s.translate(_LATIN1_MAP).encode("latin-1", "replace").decode("latin-1")

    def f(self, style: str = "", size: float = 10):
        self.set_font(self.family, style, size)


def _build_pdf(resume: dict, size: float) -> _PDF:
    pdf = _PDF()
    pdf.set_margins(14, 11, 14)
    pdf.set_auto_page_break(True, 11)
    pdf.add_page()
    lh = size * 0.47
    c = resume.get("contact") or {}
    if c.get("name"):  # document properties: some ATS/recruiter tools show these
        pdf.set_title(f"{c['name']} - Resume")
        pdf.set_author(c["name"])

    pdf.f("B", size + 8)
    pdf.cell(0, (size + 8) * 0.5, pdf.t(c.get("name", "")), align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.f("", size - 0.5)
    pdf.multi_cell(0, lh, pdf.t(_contact_line(c)), align="C", new_x="LMARGIN", new_y="NEXT")

    def heading(text: str):
        pdf.ln(lh * 0.5)
        pdf.f("B", size + 1.5)
        pdf.cell(0, lh * 1.2, pdf.t(text.upper()), new_x="LMARGIN", new_y="NEXT")
        y = pdf.get_y()
        pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
        pdf.ln(lh * 0.3)

    def bullet(text: str):
        pdf.f("", size)
        pdf.set_x(pdf.l_margin + 3)
        pdf.cell(4, lh, pdf.t("\u2022" if pdf.unicode else "-"))
        pdf.multi_cell(0, lh, pdf.t(text), new_x="LMARGIN", new_y="NEXT")

    if resume.get("summary"):
        heading("Summary")
        pdf.f("", size)
        pdf.multi_cell(0, lh, pdf.t(resume["summary"]), new_x="LMARGIN", new_y="NEXT")

    for key, title in section_order(resume):
        entries = resume.get(key) or []
        if not entries:
            continue
        heading(title)
        for e in entries:
            left, right, sub = _entry_header(key, e)
            pdf.f("B", size + 0.5)
            right_w = pdf.get_string_width(pdf.t(right)) + 2 if right else 0
            pdf.cell(pdf.epw - right_w, lh * 1.1, pdf.t(left))
            pdf.f("", size)
            pdf.cell(right_w, lh * 1.1, pdf.t(right), align="R", new_x="LMARGIN", new_y="NEXT")
            if sub:
                pdf.f("I", size)
                pdf.multi_cell(0, lh, pdf.t(sub), new_x="LMARGIN", new_y="NEXT")
            for b in e.get("bullets") or e.get("details") or []:
                bullet(b)
            pdf.ln(lh * 0.25)

    if resume.get("skills"):
        heading("Skills")
        for g in resume["skills"]:
            pdf.f("B", size)
            label = f"{g.get('category', 'Skills')}: "
            pdf.cell(pdf.get_string_width(pdf.t(label)) + 1, lh, pdf.t(label))
            pdf.f("", size)
            pdf.multi_cell(0, lh, pdf.t(", ".join(g.get("items", []))), new_x="LMARGIN", new_y="NEXT")

    for key, title in (("certifications", "Certifications"), ("awards", "Awards")):
        if resume.get(key):
            heading(title)
            for item in resume[key]:
                bullet(item)
    return pdf


def render_pdf(resume: dict, path: str | Path) -> Path:
    """Render, shrinking the font until the resume fits on one page (down to 8.5pt)."""
    path = Path(path)
    size = 10.5
    pdf = _build_pdf(resume, size)
    while pdf.pages_count > 1 and size > 8.5:
        size -= 0.5
        pdf = _build_pdf(resume, size)
    pdf.output(str(path))
    return path


def render_docx(resume: dict, path: str | Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    path = Path(path)
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Pt(36)
        s.left_margin = s.right_margin = Pt(42)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(0)

    c = resume.get("contact") or {}
    if c.get("name"):
        doc.core_properties.title = f"{c['name']} - Resume"
        doc.core_properties.author = c["name"]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(c.get("name", ""))
    r.bold, r.font.size = True, Pt(18)
    p = doc.add_paragraph(_contact_line(c))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def heading(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        run = h.add_run(text.upper())
        run.bold, run.font.size = True, Pt(12)

    if resume.get("summary"):
        heading("Summary")
        doc.add_paragraph(resume["summary"])
    for key, title in section_order(resume):
        entries = resume.get(key) or []
        if not entries:
            continue
        heading(title)
        for e in entries:
            left, right, sub = _entry_header(key, e)
            para = doc.add_paragraph()
            para.add_run(left).bold = True
            if right:
                para.add_run(f"    {right}")
            if sub:
                doc.add_paragraph().add_run(sub).italic = True
            for b in e.get("bullets") or e.get("details") or []:
                doc.add_paragraph(b, style="List Bullet")
    if resume.get("skills"):
        heading("Skills")
        for g in resume["skills"]:
            para = doc.add_paragraph()
            para.add_run(f"{g.get('category', 'Skills')}: ").bold = True
            para.add_run(", ".join(g.get("items", [])))
    for key, title in (("certifications", "Certifications"), ("awards", "Awards")):
        if resume.get(key):
            heading(title)
            for item in resume[key]:
                doc.add_paragraph(item, style="List Bullet")
    doc.save(str(path))
    return path


def render_markdown(resume: dict) -> str:
    c = resume.get("contact") or {}
    out = [f"# {c.get('name', '')}", _contact_line(c), ""]
    if resume.get("summary"):
        out += ["## Summary", resume["summary"], ""]
    for key, title in SECTIONS:
        entries = resume.get(key) or []
        if not entries:
            continue
        out.append(f"## {title}")
        for e in entries:
            left, right, sub = _entry_header(key, e)
            out.append(f"**{left}**" + (f" — {right}" if right else ""))
            if sub:
                out.append(f"*{sub}*")
            out += [f"- {b}" for b in e.get("bullets") or e.get("details") or []]
            out.append("")
    if resume.get("skills"):
        out.append("## Skills")
        out += [f"- **{g.get('category')}:** {', '.join(g.get('items', []))}" for g in resume["skills"]]
        out.append("")
    for key, title in (("certifications", "Certifications"), ("awards", "Awards")):
        if resume.get(key):
            out += [f"## {title}", *[f"- {x}" for x in resume[key]], ""]
    return "\n".join(out).strip() + "\n"


def render_letter_pdf(text: str, contact: dict, path: str | Path) -> Path:
    path = Path(path)
    pdf = _PDF()
    pdf.set_margins(22, 20, 22)
    pdf.add_page()
    pdf.f("B", 14)
    pdf.cell(0, 7, pdf.t(contact.get("name", "")), new_x="LMARGIN", new_y="NEXT")
    pdf.f("", 9.5)
    pdf.multi_cell(0, 5, pdf.t(_contact_line(contact)), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)
    pdf.f("", 11)
    pdf.multi_cell(0, 5.6, pdf.t(text), new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))
    return path
